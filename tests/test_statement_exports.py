import csv
import io
from datetime import datetime
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from api.app import app
from api.database import Base, get_db
from api.index import current_user
from api.models import AppSetting, Transaction, User, Wallet, WalletShare
from api.statements import as_csv, as_xlsx
from api import statements


@pytest.fixture
def workspace():
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        owner = User(username='Owner', email='owner@statement.test', password_hash='unused')
        viewer = User(username='Viewer', email='viewer@statement.test', password_hash='unused')
        stranger = User(username='Stranger', email='stranger@statement.test', password_hash='unused')
        db.add_all([owner, viewer, stranger]); db.flush()
        wallet = Wallet(user_id=owner.id, name='Household', initial_balance=0)
        db.add(wallet); db.flush()
        db.add_all([
            Transaction(user_id=owner.id, wallet_id=wallet.id, type='income', amount=Decimal('50.000'), description='Opening balance', date=datetime(2026, 1, 1), is_opening_balance=True),
            Transaction(user_id=owner.id, wallet_id=wallet.id, type='expense', amount=Decimal('2.250'), description='Breakfast', notes='Cafe', date=datetime(2026, 9, 2, 8, 15)),
        ])
        db.add(WalletShare(wallet_id=wallet.id, owner_id=owner.id, member_user_id=viewer.id,
                           invitee_email=viewer.email, permission='view'))
        db.commit()
        actor = {'user': viewer}
        overrides = app.dependency_overrides.copy()
        ready = getattr(app.state, 'storage_ready', False)
        app.dependency_overrides[get_db] = lambda: db
        app.dependency_overrides[current_user] = lambda: actor['user']
        app.state.storage_ready = True
        try:
            yield TestClient(app), actor, owner, viewer, stranger, wallet
        finally:
            app.dependency_overrides = overrides
            app.state.storage_ready = ready
    engine.dispose()


def params(wallet):
    return {'from_date': '2026-09-01', 'to_date': '2026-09-30', 'wallet_id': wallet.id}


def test_viewer_can_export_statement_with_prior_opening_balance(workspace):
    client, actor, owner, viewer, stranger, wallet = workspace
    db = app.dependency_overrides[get_db]()
    db.add(AppSetting(user_id=owner.id, key='currency', value='USD'))
    db.add(AppSetting(user_id=viewer.id, key='currency', value='KWD'))
    db.commit()
    choices = client.get('/api/backup/statement/wallets').json()
    assert choices == [{'id': wallet.id, 'name': 'Household', 'owner': 'Owner', 'shared': True, 'viewer': True}]
    response = client.get('/api/backup/statement', params={**params(wallet), 'format': 'csv'})
    assert response.status_code == 200, response.text
    rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert rows[0]['Entry'] == 'Opening balance'
    assert rows[0]['Balance'] == '50.000'
    assert rows[1]['Description'] == 'Breakfast'
    assert rows[1]['Money out'] == '2.250'
    assert rows[1]['Balance'] == '47.750'
    assert rows[1]['Notes'] == 'Cafe'
    assert rows[1]['Currency'] == 'USD'
    actor['user'] = stranger
    assert client.get('/api/backup/statement', params=params(wallet)).status_code == 403
    assert client.get('/api/backup/statement/wallets').json() == []


def test_xlsx_uses_real_numeric_columns_and_pdf_renders(workspace):
    client, _, _, _, _, wallet = workspace
    response = client.get('/api/backup/statement', params={**params(wallet), 'format': 'xlsx'})
    assert response.status_code == 200, response.text
    sheet = load_workbook(io.BytesIO(response.content), read_only=True).active
    assert sheet['A1'].value == 'Date'
    assert sheet['J3'].value == 47.75
    assert sheet['H3'].value == 2.25
    pdf = client.get('/api/backup/statement', params={**params(wallet), 'format': 'pdf'})
    assert pdf.status_code == 200 and pdf.content.startswith(b'%PDF')
    word = client.get('/api/backup/statement', params={**params(wallet), 'format': 'docx'})
    assert word.status_code == 200, word.text if word.status_code != 200 else ''
    document = Document(io.BytesIO(word.content))
    assert document.tables[0].rows[2].cells[5].text == 'Breakfast'


def test_statement_date_validation_does_not_change_backup(workspace):
    client, actor, owner, _, _, wallet = workspace
    invalid = client.get('/api/backup/statement', params={'from_date': '2026-09-30', 'to_date': '2026-09-01'})
    assert invalid.status_code == 422
    actor['user'] = owner
    backup = client.get('/api/backup').json()
    assert backup['version'] == 1
    assert len(backup['transactions']) == 2


def test_transfer_appears_as_money_out_and_money_in_for_accessible_wallets(workspace):
    client, actor, owner, _, _, shared = workspace
    actor['user'] = owner
    second = client.post('/api/wallets', json={'name': 'Savings', 'initial_balance': 0}).json()
    response = client.post('/api/transactions', json={
        'wallet_id': shared.id, 'transfer_wallet_id': second['id'], 'type': 'transfer',
        'amount': 10, 'description': 'Move funds', 'date': '2026-09-03T09:00:00',
    })
    assert response.status_code == 201, response.text
    exported = client.get('/api/backup/statement', params={'from_date': '2026-09-01', 'to_date': '2026-09-30'})
    assert exported.status_code == 200, exported.text
    rows = list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig'))))
    transfer = [row for row in rows if row['Description'] == 'Move funds']
    assert len(transfer) == 2
    assert {row['Entry'] for row in transfer} == {'Transfer in', 'Transfer out'}
    assert next(row for row in transfer if row['Wallet'] == 'Household')['Money out'] == '10.000'
    assert next(row for row in transfer if row['Wallet'] == 'Savings')['Money in'] == '10.000'


def test_user_text_is_not_spreadsheet_formula():
    row = ['2026-09-01', '', '=HYPERLINK("bad")', 'Owner', 'Expense',
           '  =1+1', '', 2.25, '', 47.75, 'KWD', '', '']
    csv_rows = list(csv.reader(io.StringIO(as_csv([row]).lstrip('\ufeff'))))
    assert csv_rows[1][2].startswith("'=")
    assert csv_rows[1][5].startswith("'  =")
    sheet = load_workbook(io.BytesIO(as_xlsx([row]))).active
    assert sheet['C2'].data_type == 's'
    assert sheet['F2'].data_type == 's'


def test_large_document_export_suggests_csv(workspace, monkeypatch):
    client, _, _, _, _, wallet = workspace
    monkeypatch.setattr(statements, 'statement_rows', lambda *args: [[]] * 3001)
    response = client.get('/api/backup/statement', params={**params(wallet), 'format': 'pdf'})
    assert response.status_code == 413
    assert 'use CSV' in response.json()['detail']
