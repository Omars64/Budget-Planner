"""Read-only wallet statements for owners and shared-wallet members."""
import csv
import io
from datetime import date, datetime, time
from decimal import Decimal
from html import escape

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import or_

from .database import get_db
from .index import current_user, materialize_recurring_for_user, normalize_email, setting
from .models import Category, Transaction, User, Wallet, WalletShare
from .account_security import audit, limit

router = APIRouter()
HEADINGS = ['Date', 'Time', 'Wallet', 'Owner', 'Entry', 'Description', 'Category',
            'Money out', 'Money in', 'Balance', 'Currency', 'Notes', 'Recorded by']


def allowed_wallets(db, user):
    shares = db.query(WalletShare.wallet_id).filter(or_(
        WalletShare.member_user_id == user.id,
        WalletShare.invitee_email == normalize_email(user.email),
        WalletShare.owner_id == user.id,
    ))
    return db.query(Wallet).filter(or_(Wallet.user_id == user.id, Wallet.id.in_(shares))).order_by(Wallet.name, Wallet.id).all()


@router.get('/api/backup/statement/wallets')
def statement_wallets(user: User = Depends(current_user), db=Depends(get_db)):
    wallets = allowed_wallets(db, user)
    owner_ids = {wallet.user_id for wallet in wallets}
    owners = {row.id: row.username for row in db.query(User).filter(User.id.in_(owner_ids)).all()} if owner_ids else {}
    shared_ids = {row[0] for row in db.query(WalletShare.wallet_id).filter(WalletShare.wallet_id.in_([w.id for w in wallets])).all()} if wallets else set()
    return [{'id': w.id, 'name': w.name, 'owner': owners.get(w.user_id, ''),
             'shared': w.id in shared_ids, 'viewer': w.user_id != user.id} for w in wallets]


def statement_rows(db, user, wallet_id, start, end):
    wallets = allowed_wallets(db, user)
    selected = [w for w in wallets if wallet_id is None or w.id == wallet_id]
    if wallet_id is not None and not selected:
        raise HTTPException(403, 'This wallet is not available to you')
    owners = {w.user_id for w in selected}
    for owner_id in owners:
        materialize_recurring_for_user(db, owner_id)
    owner_names = {u.id: u.username for u in db.query(User).filter(User.id.in_(owners)).all()} if owners else {}
    ids = [w.id for w in selected]
    category_ids = set()
    txs = db.query(Transaction).filter(or_(Transaction.wallet_id.in_(ids), Transaction.transfer_wallet_id.in_(ids)),
                                       Transaction.date <= datetime.combine(end, time.max)).order_by(Transaction.date, Transaction.id).limit(100001).all() if ids else []
    if len(txs) > 100000:
        raise HTTPException(413, 'This statement has too many records. Choose a shorter date range or one wallet.')
    category_ids.update(t.category_id for t in txs if t.category_id)
    categories = {c.id: c.name for c in db.query(Category).filter(Category.id.in_(category_ids)).all()} if category_ids else {}
    recorder_ids = {t.recorded_by_id for t in txs if t.recorded_by_id}
    recorders = {u.id: u.username for u in db.query(User).filter(User.id.in_(recorder_ids)).all()} if recorder_ids else {}
    rows = []
    by_wallet = {w.id: [] for w in selected}
    for tx in txs:
        if tx.wallet_id in by_wallet:
            by_wallet[tx.wallet_id].append((tx, False))
        if tx.type == 'transfer' and tx.transfer_wallet_id in by_wallet:
            by_wallet[tx.transfer_wallet_id].append((tx, True))
    for wallet in selected:
        currency = setting(db, wallet.user_id, 'currency', 'KWD')
        balance = Decimal(str(wallet.initial_balance or 0))
        entries = []
        for tx, incoming in by_wallet[wallet.id]:
            outgoing = not incoming
            amount = Decimal(str(tx.amount))
            money_in = amount if incoming or (outgoing and tx.type == 'income') else Decimal('0')
            money_out = amount if outgoing and tx.type in {'expense', 'transfer'} else Decimal('0')
            if tx.date.date() < start:
                balance += money_in - money_out
                continue
            entries.append((tx, money_out, money_in, incoming))
        rows.append([start.isoformat(), '', wallet.name, owner_names.get(wallet.user_id, ''),
                     'Opening balance', 'Balance before selected period', '', '', '',
                     float(balance), currency, '', ''])
        for tx, money_out, money_in, incoming in entries:
            balance += money_in - money_out
            entry = 'Opening funds' if tx.is_opening_balance else ('Transfer in' if incoming else 'Transfer out' if tx.type == 'transfer' else tx.type.title())
            rows.append([tx.date.date().isoformat(), tx.date.strftime('%H:%M'), wallet.name,
                         owner_names.get(wallet.user_id, ''), entry, tx.description,
                         categories.get(tx.category_id, ''), float(money_out) if money_out else '',
                         float(money_in) if money_in else '', float(balance), currency,
                         tx.notes or '', recorders.get(tx.recorded_by_id, '')])
    return rows


def as_csv(rows):
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(HEADINGS)
    writer.writerows(csv_row(row) for row in rows)
    return '\ufeff' + output.getvalue()


def csv_row(row):
    values = display_row(row)
    for index, value in enumerate(values):
        if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')) and index not in (7, 8, 9):
            values[index] = "'" + value
    return values


def display_row(row):
    values = list(row)
    for index in (7, 8, 9):
        if values[index] != '':
            values[index] = f'{values[index]:,.3f}'
    return values


def as_xlsx(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Statement'
    ws.append(HEADINGS)
    for row in rows:
        ws.append([None if value == '' else value for value in row])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17324A')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    widths = [13, 9, 22, 20, 19, 34, 20, 15, 15, 17, 12, 38, 20]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for index in (7, 8, 9):
            row[index].number_format = '#,##0.000;[Red](#,##0.000)'
        if row[4].value == 'Opening balance':
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='EAF1F5')
                cell.font = Font(bold=True)
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                cell.data_type = 's'
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def as_pdf(rows, title):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A3
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A3), leftMargin=22, rightMargin=22, topMargin=26, bottomMargin=24)
    styles = getSampleStyleSheet()
    visible = list(range(len(HEADINGS)))
    data = [[HEADINGS[i] for i in visible]]
    for row in rows:
        data.append([Paragraph(escape(str(display_row(row)[i])) if row[i] != '' else '', styles['BodyText']) for i in visible])
    table = Table(data, colWidths=[67, 38, 85, 77, 83, 135, 85, 69, 69, 72, 53, 170, 76], repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17324A')),
                               ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                               ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6F8')]),
                               ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('FONTSIZE', (0, 0), (-1, -1), 8),
                               ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
    doc.build([Paragraph(escape(title), styles['Title']), Spacer(1, 10), table])
    return output.getvalue()


def as_docx(rows, title):
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(16.54), Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.55)
    doc.add_heading(title, 0)
    doc.add_paragraph('Amounts are shown in the currency column. Opening balance is the balance before the selected period.')
    visible = list(range(len(HEADINGS)))
    table = doc.add_table(rows=1, cols=len(visible))
    table.style = 'Light Shading Accent 1'
    for cell, i in zip(table.rows[0].cells, visible):
        cell.text = HEADINGS[i]
    for row in rows:
        values = display_row(row)
        for cell, i in zip(table.add_row().cells, visible):
            cell.text = str(values[i]) if row[i] != '' else ''
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


@router.get('/api/backup/statement')
def download_statement(format: str = Query('csv', pattern='^(csv|xlsx|pdf|docx)$'),
                       from_date: date = Query(...), to_date: date = Query(...),
                       wallet_id: int | None = Query(None, ge=1),
                       user: User = Depends(current_user), db=Depends(get_db)):
    if from_date > to_date:
        raise HTTPException(422, 'From date must be on or before To date')
    if (to_date - from_date).days > 3650:
        raise HTTPException(422, 'Choose a range of ten years or less')
    limit(db, f'statement:{user.id}', 30, 3600)
    rows = statement_rows(db, user, wallet_id, from_date, to_date)
    row_limit = {'pdf': 3000, 'docx': 3000, 'xlsx': 20000}.get(format)
    if row_limit and len(rows) > row_limit:
        raise HTTPException(413, f'This format supports up to {row_limit:,} rows. Choose one wallet or a shorter range, or use CSV.')
    title = f'Budgetly statement | {from_date.isoformat()} to {to_date.isoformat()}'
    renderers = {'csv': (as_csv, 'text/csv; charset=utf-8'),
                 'xlsx': (as_xlsx, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                 'pdf': (lambda r: as_pdf(r, title), 'application/pdf'),
                 'docx': (lambda r: as_docx(r, title), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}
    renderer, media_type = renderers[format]
    content = renderer(rows)
    audit(db, user.id, user.id, 'Exported wallet statement', f'wallet:{wallet_id}' if wallet_id else 'all-wallets')
    db.commit()
    return Response(content=content, media_type=media_type,
                    headers={'Content-Disposition': f'attachment; filename="budgetly-statement-{from_date}-{to_date}.{format}"',
                             'Cache-Control': 'no-store'})
